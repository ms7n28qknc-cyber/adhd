#!/usr/bin/env python3
"""
extract_hscims.py
=================
Reads the HSCIMS 2025 data-tables workbook and extracts mental health
indicators for each HSC Trust and Northern Ireland overall.

The workbook structure (per sheet per indicator):
  Row N:   Indicator name (string)
  Row N+1: Unit label | period_1 | period_2 | ... | period_5 | 'Trend Analysis'
  Row N+2: 'Northern Ireland' | v1 | v2 | v3 | v4 | v5 | ...
  Row N+3: [Trust/Deprivation rows]

NI sheet rows:
  'Northern Ireland', '1 (Most Deprived)', '5 (Least Deprived)'

Trust sheet rows (e.g., Belfast HSCT):
  'Northern Ireland', 'Belfast Trust', 'Belfast Trust Deprived'

Usage
-----
    python extract_hscims.py                    # auto-detect file, write to site/data/
    python extract_hscims.py --explore          # dump sheet structure
    python extract_hscims.py --xlsx /path/to/file.xlsx --out /path/to/out.json
"""

import sys
import os
import re
import json
import openpyxl

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

XLSX_SEARCH_PATHS = [
    "site/data/hscims-report-2025-data-tables-by-area.xlsx",
    "data/hscims-report-2025-data-tables-by-area.xlsx",
    "/data/hscims-report-2025-data-tables-by-area.xlsx",
]

OUT_DEFAULT = "site/data/health-context.json"

SHEET_NI = "Northern Ireland"
SHEET_TRUSTS = {
    "Belfast":       "Belfast HSCT",
    "Northern":      "Northern HSCT",
    "South Eastern": "South Eastern HSCT",
    "Southern":      "Southern HSCT",
    "Western":       "Western HSCT",
}

# Indicator keyword → output key, scale (multiply by to convert to per 1,000)
INDICATORS = {
    "moodAnxiety": {
        "keyword":   "Mood & Anxiety",
        "scale":     1.0,       # already per 1,000 population
        "required":  True,
    },
    "selfHarm": {
        "keyword":   "Self-Harm",
        "scale":     1 / 100,   # per 100,000 → per 1,000
        "required":  True,
    },
    "suicide": {
        "keyword":   "Suicide",
        "scale":     1 / 100,
        "required":  True,
    },
    "drugAdmissions": {
        "keyword":   "Drug Related Causes",
        "scale":     1 / 100,
        "required":  False,
    },
    "alcoholDeaths": {
        "keyword":   "Alcohol Specific",
        "scale":     1 / 100,
        "required":  False,
    },
}

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def find_xlsx(override=None):
    if override:
        if os.path.isfile(override):
            return override
        raise FileNotFoundError(f"Excel file not found: {override}")
    for p in XLSX_SEARCH_PATHS:
        if os.path.isfile(p):
            return p
    raise FileNotFoundError(
        "Excel file not found. Tried:\n  " + "\n  ".join(XLSX_SEARCH_PATHS)
    )


def period_end_year(label):
    """
    Extract the end year from a period label as a string.

    Examples:
      2019              → '2019'  (int or str integer)
      '2017-19'         → '2019'
      '2019-23'         → '2023'
      '2015/16-2019/20' → '2020'
      '2019/20-2023/24' → '2024'
    """
    if isinstance(label, (int, float)):
        return str(int(label))
    s = str(label).strip()
    # Financial year range: '2019/20-2023/24'
    m = re.search(r'(\d{4})/(\d{2})\s*$', s)
    if m:
        return str(int(m.group(1)) + 1)
    # Calendar year range ending in 2-digit year: '2017-19', '2019-23'
    m = re.search(r'(\d{4})-(\d{2})\s*$', s)
    if m:
        century = m.group(1)[:2]
        return century + m.group(2)
    # Plain 4-digit year: '2023'
    m = re.match(r'^(\d{4})$', s)
    if m:
        return m.group(1)
    return s


def format_period_label(label):
    """Return a clean display label, replacing '-' with '–' (en-dash)."""
    if isinstance(label, (int, float)):
        return str(int(label))
    return str(label).strip().replace('-', '–')


def cell_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', '').strip())
    except (ValueError, TypeError):
        return None


def row_vals(row, pad_to=8):
    lst = list(row)
    while len(lst) < pad_to:
        lst.append(None)
    return lst


# ---------------------------------------------------------------------------
# Sheet exploration
# ---------------------------------------------------------------------------

def explore_sheet(ws, max_rows=80):
    print(f"\n{'='*60}\nSheet: {ws.title}\n{'='*60}")
    for i, row in enumerate(ws.iter_rows(max_row=max_rows, values_only=True), 1):
        non_empty = [(j, v) for j, v in enumerate(row) if v is not None]
        if non_empty:
            vals = '  |  '.join(f'[{j}] {repr(v)[:50]}' for j, v in non_empty[:8])
            print(f'  row {i:3d}: {vals}')


def explore(wb):
    for sheet_name in [SHEET_NI] + list(SHEET_TRUSTS.values()):
        if sheet_name in wb.sheetnames:
            explore_sheet(wb[sheet_name])
        else:
            print(f"\n[Sheet not found: {sheet_name}]")

# ---------------------------------------------------------------------------
# Core extraction
# ---------------------------------------------------------------------------

def find_indicator_row(rows, keyword, max_rows=None):
    """Return 0-based row index of first row whose col-0 contains keyword."""
    kw_lower = keyword.lower()
    limit = max_rows or len(rows)
    for i, row in enumerate(rows[:limit]):
        v = row[0] if row else None
        if isinstance(v, str) and kw_lower in v.lower():
            return i
    return None


def extract_ni_block(rows, start_idx, scale):
    """
    Extract data from NI sheet indicator block.
    Returns:
      ni_vals          : {year_key: scaled_float}
      q1_vals          : {year_key: scaled_float}   (most deprived)
      q5_vals          : {year_key: scaled_float}   (least deprived)
      period_labels    : {year_key: display_label}
      year_keys        : [ordered list of year_key strings]
    """
    # Row start_idx = indicator name; start_idx+1 = column headers
    header_row = rows[start_idx + 1] if start_idx + 1 < len(rows) else []
    periods = [v for v in header_row if v is not None and str(v) != 'Trend Analysis'][1:]
    # periods[0..4] correspond to columns 1..5

    year_keys     = [period_end_year(p) for p in periods]
    period_labels = {yk: format_period_label(p) for yk, p in zip(year_keys, periods)}

    ni_vals = {}
    q1_vals = {}
    q5_vals = {}

    # Scan rows below header for the data rows we need
    for row in rows[start_idx + 2: start_idx + 8]:
        label = row[0] if row else None
        if not isinstance(label, str):
            continue
        label_l = label.lower()
        vals = [cell_float(row[j]) for j in range(1, 6)]

        def scaled(v):
            return round(v * scale, 4) if v is not None else None

        if label_l == 'northern ireland':
            ni_vals = {yk: scaled(v) for yk, v in zip(year_keys, vals)}
        elif '1 (most deprived)' in label_l:
            q1_vals = {yk: scaled(v) for yk, v in zip(year_keys, vals)}
        elif '5 (least deprived)' in label_l:
            q5_vals = {yk: scaled(v) for yk, v in zip(year_keys, vals)}

    return ni_vals, q1_vals, q5_vals, period_labels, year_keys


def extract_trust_block(rows, start_idx, scale, trust_name):
    """
    Extract Trust average and Trust Deprived from a Trust sheet indicator block.
    Returns:
      avg_vals         : {year_key: scaled_float}
      deprived_vals    : {year_key: scaled_float}
      year_keys        : [ordered list]
    """
    header_row = rows[start_idx + 1] if start_idx + 1 < len(rows) else []
    periods  = [v for v in header_row if v is not None and str(v) != 'Trend Analysis'][1:]
    year_keys = [period_end_year(p) for p in periods]

    avg_vals     = {}
    deprived_vals = {}
    # Trust rows are labeled "[Trust Name] Trust" and "[Trust Name] Trust Deprived"
    # e.g. 'Belfast Trust', 'Belfast Trust Deprived'
    # Match by looking for trust_name in the row label (case-insensitive)
    tn_lower = trust_name.lower()

    for row in rows[start_idx + 2: start_idx + 10]:
        label = row[0] if row else None
        if not isinstance(label, str):
            continue
        label_l = label.lower()
        if tn_lower not in label_l:
            continue
        vals = [cell_float(row[j]) for j in range(1, 6)]

        def scaled(v):
            return round(v * scale, 4) if v is not None else None

        if label_l == f"{tn_lower} trust deprived":
            deprived_vals = {yk: scaled(v) for yk, v in zip(year_keys, vals)}
        elif label_l == f"{tn_lower} trust":
            avg_vals = {yk: scaled(v) for yk, v in zip(year_keys, vals)}

    return avg_vals, deprived_vals, year_keys

# ---------------------------------------------------------------------------
# Main extraction
# ---------------------------------------------------------------------------

def extract_all(wb):
    output = {}

    # Load sheets
    ni_sheet = wb[SHEET_NI] if SHEET_NI in wb.sheetnames else None
    if ni_sheet is None:
        print(f"ERROR: '{SHEET_NI}' sheet not found.", file=sys.stderr)
        return {}
    ni_rows_raw = list(ni_sheet.iter_rows(values_only=True))
    ni_rows     = [row_vals(r) for r in ni_rows_raw]

    trust_rows = {}
    for trust_name, sheet_name in SHEET_TRUSTS.items():
        if sheet_name not in wb.sheetnames:
            print(f"WARNING: Sheet '{sheet_name}' not found.", file=sys.stderr)
            continue
        ws = wb[sheet_name]
        trust_rows[trust_name] = [row_vals(r) for r in ws.iter_rows(values_only=True)]

    for ind_key, ind_cfg in INDICATORS.items():
        keyword = ind_cfg["keyword"]
        scale   = ind_cfg["scale"]
        print(f"\nExtracting: {ind_key}  (keyword: '{keyword}')")

        # ── NI sheet ──
        ni_idx = find_indicator_row(ni_rows, keyword)
        if ni_idx is None:
            msg = f"  Indicator not found in NI sheet"
            print(msg)
            if ind_cfg["required"]:
                print(f"  Run with --explore to inspect the sheet structure.")
            continue

        ni_vals, q1_vals, q5_vals, period_labels, year_keys = \
            extract_ni_block(ni_rows, ni_idx, scale)
        print(f"  NI sheet: row {ni_idx+1}, {len(ni_vals)} values, years: {year_keys}")

        # ── Trust sheets ──
        trusts_out = {}
        for trust_name, rows in trust_rows.items():
            t_idx = find_indicator_row(rows, keyword)
            if t_idx is None:
                print(f"  {trust_name}: not found in sheet")
                continue
            avg_vals, deprived_vals, t_year_keys = \
                extract_trust_block(rows, t_idx, scale, trust_name)
            if avg_vals:
                trusts_out[trust_name] = {
                    "average":      avg_vals,
                    "mostDeprived": deprived_vals,
                }
                print(f"  {trust_name}: {len(avg_vals)} values")
            else:
                print(f"  {trust_name}: data rows not matched (check trust row labels)")

        output[ind_key] = {
            "years":        year_keys,
            "periodLabels": period_labels,
            "ni":           ni_vals,
            "trusts":       trusts_out,
            "deprivationQuintiles": {
                "mostDeprived":  q1_vals,
                "leastDeprived": q5_vals,
            },
        }

    return output

# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    args      = sys.argv[1:]
    explore_mode = "--explore" in args
    xlsx_path = None
    out_path  = OUT_DEFAULT

    i = 0
    while i < len(args):
        if args[i] == "--xlsx" and i + 1 < len(args):
            xlsx_path = args[i + 1]; i += 2
        elif args[i] == "--out" and i + 1 < len(args):
            out_path = args[i + 1]; i += 2
        else:
            i += 1

    try:
        xlsx_file = find_xlsx(xlsx_path)
    except FileNotFoundError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)

    print(f"Loading: {xlsx_file}")
    wb = openpyxl.load_workbook(xlsx_file, data_only=True)
    print(f"Sheets:  {wb.sheetnames}")

    if explore_mode:
        explore(wb)
        return

    data = extract_all(wb)
    if not data:
        print("\nERROR: Nothing extracted. Run with --explore to inspect structure.")
        sys.exit(1)

    out_dir = os.path.dirname(out_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, separators=(",", ":"), indent=2)

    print(f"\nWritten: {out_path}")
    print(f"Indicators: {list(data.keys())}")


if __name__ == "__main__":
    main()
